"""
Excel I/O helpers for the web platform runner.
Extracted from run_tests.py and adapted to accept a results dict parameter.
"""
import shutil
from typing import List, Dict, Any

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# Excel column names expected in the input file
COL_ID       = "Test Case #"
COL_TITLE    = "Title"
COL_STEPS    = "Steps"
COL_EXPECTED = "Expected Result"
COL_STATUS   = "Status"
COL_ACTUAL   = "Actual Result"
COL_COMMENTS = "Comments"

STATUS_LABELS = {
    "PASS":   "✅ PASS",
    "FAIL":   "❌ FAIL",
    "MANUAL": "⚠️ MANUAL",
    "SKIP":   "— SKIP",
}

FILL = {
    "PASS":   PatternFill("solid", fgColor="C6EFCE"),
    "FAIL":   PatternFill("solid", fgColor="FFC7CE"),
    "MANUAL": PatternFill("solid", fgColor="FFEB9C"),
    "SKIP":   PatternFill("solid", fgColor="D9D9D9"),
}


def _col_index(headers: dict, *names: str):
    for name in names:
        if name in headers:
            return headers[name]
    return None


def load_test_cases(path: str) -> List[Dict[str, Any]]:
    """
    Load test-case definitions from an Excel file.
    Returns a list of dicts: {id, title, steps, expected, _row}.
    Raises ValueError if the file appears empty or malformed.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active

    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val:
            headers[str(val).strip()] = c

    tc_col  = _col_index(headers, COL_ID,       "Test Case Number", "TC #", "ID")
    ttl_col = _col_index(headers, COL_TITLE,     "Test Name", "Name")
    stp_col = _col_index(headers, COL_STEPS,     "Test Steps", "Step")
    exp_col = _col_index(headers, COL_EXPECTED,  "Expected", "Expected Outcome")

    if not tc_col:
        raise ValueError(
            f"Could not find a 'Test Case #' column. Found: {list(headers.keys())}"
        )

    test_cases = []
    for r in range(2, ws.max_row + 1):
        tc_id = ws.cell(r, tc_col).value
        if not tc_id:
            continue
        test_cases.append({
            "id":       str(tc_id).strip(),
            "title":    ws.cell(r, ttl_col).value if ttl_col else "",
            "steps":    ws.cell(r, stp_col).value if stp_col else "",
            "expected": ws.cell(r, exp_col).value if exp_col else "",
            "_row":     r,
        })

    wb.close()

    if not test_cases:
        raise ValueError(f"No test-case rows found in {path}.")

    return test_cases


def write_results(
    input_path: str,
    output_path: str,
    results: Dict[str, Dict[str, str]],
) -> None:
    """
    Copy input_path → output_path, then fill in Status / Actual Result / Comments
    from the *results* dict: {test_case_id: {status, actual, comment}}.
    """
    shutil.copy(input_path, output_path)
    wb = load_workbook(output_path)
    ws = wb.active

    headers: Dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val:
            headers[str(val).strip()] = c

    tc_col = _col_index(headers, COL_ID,       "Test Case Number", "TC #", "ID")
    st_col = _col_index(headers, COL_STATUS,   "Status")
    ar_col = _col_index(headers, COL_ACTUAL,   "Actual Result")
    cm_col = _col_index(headers, COL_COMMENTS, "Comments")

    # Add missing output columns
    next_col = ws.max_column + 1
    if not st_col:
        ws.cell(1, next_col).value = COL_STATUS
        st_col = next_col; next_col += 1
    if not ar_col:
        ws.cell(1, next_col).value = COL_ACTUAL
        ar_col = next_col; next_col += 1
    if not cm_col:
        ws.cell(1, next_col).value = COL_COMMENTS
        cm_col = next_col

    for r in range(2, ws.max_row + 1):
        tc_id = ws.cell(r, tc_col).value if tc_col else None
        if not tc_id:
            continue
        tc_id = str(tc_id).strip()
        if tc_id not in results:
            continue

        result = results[tc_id]
        status = result["status"]
        label  = STATUS_LABELS.get(status, status)
        fill   = FILL.get(status)

        for col_idx, value in [
            (st_col, label),
            (ar_col, result.get("actual", "")),
            (cm_col, result.get("comment", "")),
        ]:
            if col_idx:
                cell = ws.cell(r, col_idx)
                cell.value = value
                if fill:
                    cell.fill = fill

    wb.save(output_path)
